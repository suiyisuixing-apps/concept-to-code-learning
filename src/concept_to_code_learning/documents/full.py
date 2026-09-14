"""Safe local parsing and storage for PDF, PPTX, DOCX and Markdown."""

from __future__ import annotations

import asyncio
import base64
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from tempfile import mkdtemp
from xml.parsers import expat

from concept_to_code_learning.full_contracts.models import (
    Block,
    DocumentRecord,
    DocumentUnit,
    Preview,
    ProviderCapability,
    SourceLocator,
)
from concept_to_code_learning.full_learning.errors import LearningError
from concept_to_code_learning.full_learning.io import run_io, settle
from concept_to_code_learning.full_learning.ports import DocumentAsset, DocumentUpload

MAX_UNPACKED = 100 * 1024 * 1024
MAX_ENTRIES = 2_000
MAX_ASSETS = 300
MAX_UNITS = 300
MAX_METADATA_BYTES = 40 * 1024 * 1024
SAFE_IMAGES = {"image/png", "image/jpeg", "image/webp"}


def fail(code, message, status=422):
    return LearningError(code, "document", message, status)


def locator(kind, index, headings, block_id=None):
    return SourceLocator(unit_type=kind, index=index, heading_path=headings, block_id=block_id)


def block(kind, index, number, block_kind, text, headings, rows=None, asset_id=None, bbox=None):
    block_id = f"{kind}-{index}-block-{number}"
    return Block(block_id=block_id, kind=block_kind, text=text,
                 source_locator=locator(kind, index, headings, block_id),
                 table_rows=rows or [], image_asset_id=asset_id, bbox=bbox)


def validate_office(content, source_type):
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
        infos = archive.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise fail("INVALID_FILE", f"{source_type} 文件包已损坏。") from exc
    if not infos or len(infos) > MAX_ENTRIES:
        archive.close()
        raise fail("INVALID_FILE", "Office 文件包含异常数量的资源。")
    total = 0
    for info in infos:
        path = PurePosixPath(info.filename.replace("\\", "/"))
        if path.is_absolute() or ".." in path.parts:
            archive.close()
            raise fail("INVALID_FILE", "Office 文件包含路径穿越条目。")
        total += info.file_size
        if total > MAX_UNPACKED or info.file_size > MAX_UNPACKED:
            archive.close()
            raise fail("FILE_TOO_LARGE", "Office 文件解包后超过安全上限。", 413)
    names = {item.filename for item in infos}
    required = "ppt/presentation.xml" if source_type == "PPTX" else "word/document.xml"
    if required not in names or "[Content_Types].xml" not in names:
        archive.close()
        raise fail("INVALID_FILE", f"扩展名与 {source_type} 结构不匹配。")
    defaults, overrides = {}, {}
    types_namespace = "http://schemas.openxmlformats.org/package/2006/content-types}"
    depth = 0

    def content_type(name, attrs):
        nonlocal depth
        depth += 1
        if depth == 1 and name != types_namespace + "Types":
            raise fail("INVALID_FILE", "Office 内容类型索引无效。")
        if depth != 2:
            return
        if name == types_namespace + "Default":
            defaults[attrs.get("Extension", "").lower()] = attrs.get("ContentType", "")
        elif name == types_namespace + "Override":
            # OPC libraries compare overrides case-insensitively and preserve
            # percent escapes in ZIP member names; match their dispatch exactly.
            overrides[attrs.get("PartName", "").lower()] = attrs.get("ContentType", "")

    def end_content_type(name):
        nonlocal depth
        depth -= 1

    def reject_entity(*args):
        raise fail("INVALID_FILE", "Office 文件包含不允许的 XML 实体。")

    def check_xml(raw, handler=None, end_handler=None):
        # Parse original bytes: XML declares its own encoding. A UTF-8 text scan
        # misses UTF-16, while extension-only checks miss OPC XML parts in .bin.
        parser = expat.ParserCreate(namespace_separator="}")
        parser.StartDoctypeDeclHandler = reject_entity
        parser.EntityDeclHandler = reject_entity
        parser.ExternalEntityRefHandler = reject_entity
        parser.StartElementHandler = handler
        parser.EndElementHandler = end_handler
        parser.Parse(raw, True)

    try:
        check_xml(archive.read("[Content_Types].xml"), content_type, end_content_type)
        for name in names - {"[Content_Types].xml"}:
            mime = overrides.get("/" + name.lower(), defaults.get(PurePosixPath(name).suffix[1:].lower(), ""))
            if name.casefold().endswith((".xml", ".rels")) or mime.casefold().endswith(("+xml", "/xml")):
                check_xml(archive.read(name))
    except (expat.ExpatError, zipfile.BadZipFile, RuntimeError) as exc:
        raise fail("INVALID_FILE", "Office 文件的 XML 无法安全读取。") from exc
    finally:
        archive.close()


class FullDocumentProvider:
    def __init__(self, settings):
        self.settings = settings
        self.root = settings.data_dir / "inogi-sama" / "documents"
        self.root.mkdir(parents=True, exist_ok=True)

    async def capabilities(self):
        missing = [name for name in ("pypdf", "pptx", "docx") if importlib.util.find_spec(name) is None]
        return ProviderCapability(
            provider_id="inogi-document", implemented=True, available=not missing, mode="LIVE",
            features=["pdf-native-pages", "pptx-learning-view", "docx-sections",
                      "markdown-safe-blocks", "unicode-codepoint-selection", "local-assets"],
            status="UNAVAILABLE" if missing else "AVAILABLE",
            reason_code="DEPENDENCY_MISSING" if missing else None,
            needed_action="安装正式依赖后重试。" if missing else "扫描 PDF 需 OCR；Office 学习视图不等同原版。")

    async def close(self):
        return None

    async def import_document(self, upload: DocumentUpload):
        content = bytes(upload.content)
        if not content:
            raise fail("INVALID_FILE", "上传文件为空。")
        if len(content) > self.settings.max_document_bytes:
            raise fail("FILE_TOO_LARGE", "文件超过导入上限。", 413)
        source_type = self.detect(upload.file_name, content)
        sha = hashlib.sha256(content).hexdigest()
        document_id = "doc-" + hashlib.sha256(content + b"\0" + upload.file_name.encode("utf-8")).hexdigest()[:24]
        final = self.root / document_id
        if final.exists():
            return await run_io(self.record, final)
        staging = Path(mkdtemp(prefix=".import-", dir=self.root))
        try:
            suffix = Path(upload.file_name).suffix.lower()
            original = f"original{suffix}"
            units, assets, warnings, capabilities = await self.parse_async(source_type, content, document_id)
            no_text = units and all(u.extraction_status == "NO_EXTRACTABLE_TEXT" for u in units)
            partial = any(u.extraction_status != "READY" for u in units)
            record = DocumentRecord(
                mode="LIVE", document_id=document_id, file_name=upload.file_name,
                source_type=source_type, original_sha256=sha, revision=1, unit_count=len(units),
                import_status="NO_EXTRACTABLE_TEXT" if no_text else ("PARTIAL" if partial else "READY"),
                capabilities=capabilities, warnings=warnings, created_at=datetime.now(timezone.utc))
            metadata = {"record": record.model_dump(mode="json"), "original": original,
                        "units": [u.model_dump(mode="json") for u in units],
                        "assets": {key: {"media_type": val.media_type, "file_name": val.file_name}
                                   for key, val in assets.items()}}
            record = await run_io(self._stage, staging, original, content, metadata, assets)
            # The expensive work is cancellable before publication. This short
            # atomic rename has no await between the decision and its result.
            try:
                os.replace(staging, final)
            except OSError:
                if not final.is_dir():
                    raise
                record = await run_io(self.record, final)
                await run_io(self._remove_staging, staging)
            return record
        except (LearningError, asyncio.CancelledError):
            await run_io(self._remove_staging, staging)
            raise
        except Exception as exc:
            await run_io(self._remove_staging, staging)
            raise fail("INVALID_FILE", "文档损坏、加密或无法安全读取。") from exc

    @staticmethod
    def _remove_staging(staging):
        if staging.is_dir():
            for path in staging.glob("original*"):
                if not path.is_symlink():
                    path.chmod(0o600)
            shutil.rmtree(staging)

    def _stage(self, staging, original, content, metadata, assets):
        encoded = json.dumps(metadata, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        if len(encoded) > MAX_METADATA_BYTES:
            raise fail("FILE_TOO_LARGE", "提取结果超过安全上限，请拆分文件后重试。", 413)
        (staging / "assets").mkdir()
        (staging / original).write_bytes(content)
        for asset_id, asset in assets.items():
            (staging / "assets" / asset_id).write_bytes(asset.content)
        (staging / "metadata.json").write_bytes(encoded)
        record = self.record(staging)
        (staging / original).chmod(0o444)
        return record

    async def list_documents(self):
        return await run_io(self._list_documents)

    def _list_documents(self):
        values = []
        for path in self.root.glob("doc-*/metadata.json"):
            try:
                values.append(self.record(path.parent))
            except (OSError, ValueError, KeyError, json.JSONDecodeError, LearningError):
                continue
        return sorted(values, key=lambda item: item.created_at, reverse=True)

    async def get_document(self, document_id):
        try:
            return await run_io(lambda: self.record(self.folder(document_id)))
        except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
            raise fail("FILE_NOT_FOUND", "文档不存在或索引无法读取。", 404) from exc

    async def list_units(self, document_id):
        return await run_io(self._list_units, document_id)

    def _list_units(self, document_id):
        data = self.metadata(self.folder(document_id))
        return [DocumentUnit.model_validate(item) for item in data["units"]]

    async def get_unit(self, document_id, unit_id):
        for unit in await self.list_units(document_id):
            if unit.unit_id == unit_id:
                return unit
        raise fail("FILE_NOT_FOUND", "文档单元不存在。", 404)

    async def get_asset(self, document_id, asset_id):
        return await run_io(self._get_asset, document_id, asset_id)

    def _get_asset(self, document_id, asset_id):
        folder = self.folder(document_id)
        data = self.metadata(folder)
        if asset_id == "original-pdf" and data["record"]["source_type"] == "PDF":
            return DocumentAsset((folder / data["original"]).read_bytes(), "application/pdf",
                                 data["record"]["file_name"])
        if not re.fullmatch(r"[A-Za-z0-9_-]+", asset_id):
            raise fail("FILE_NOT_FOUND", "文档资产不存在。", 404)
        info = data.get("assets", {}).get(asset_id)
        if not info:
            raise fail("FILE_NOT_FOUND", "文档资产不存在。", 404)
        path = folder / "assets" / asset_id
        if not path.is_file() or path.is_symlink():
            raise fail("FILE_NOT_FOUND", "文档资产不存在。", 404)
        return DocumentAsset(path.read_bytes(), info["media_type"], info["file_name"])

    def folder(self, document_id):
        if not re.fullmatch(r"doc-[a-f0-9]{24}", document_id):
            raise fail("FILE_NOT_FOUND", "文档不存在。", 404)
        folder = self.root / document_id
        if folder.is_symlink():
            raise fail("FILE_NOT_FOUND", "文档路径不合法。", 404)
        return folder

    @staticmethod
    def metadata(folder):
        try:
            path = folder / "metadata.json"
            if path.is_symlink() or path.stat().st_size > MAX_METADATA_BYTES:
                raise fail("DOCUMENT_VERSION_MISMATCH", "文档索引无效。", 409)
            data = json.loads(path.read_text(encoding="utf-8"))
            original = folder / data["original"]
            if (Path(data["original"]).name != data["original"] or original.is_symlink()
                    or hashlib.sha256(original.read_bytes()).hexdigest() != data["record"]["original_sha256"]):
                raise fail("DOCUMENT_VERSION_MISMATCH", "文档副本已变化，请重新导入。", 409)
            return data
        except (OSError, json.JSONDecodeError) as exc:
            raise fail("FILE_NOT_FOUND", "文档不存在或索引无法读取。", 404) from exc

    def record(self, folder):
        return DocumentRecord.model_validate(self.metadata(folder)["record"])

    @staticmethod
    def detect(file_name, content):
        suffix = Path(file_name).suffix.lower()
        if suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise fail("INVALID_FILE", "扩展名与 PDF 结构不匹配。")
            return "PDF"
        if suffix in {".pptx", ".docx"}:
            if not content.startswith(b"PK"):
                raise fail("INVALID_FILE", "扩展名与 Office 结构不匹配。")
            return suffix[1:].upper()
        if suffix in {".md", ".markdown"}:
            if b"\x00" in content:
                raise fail("INVALID_FILE", "Markdown 必须是 UTF-8 文本。")
            try:
                content.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise fail("INVALID_FILE", "Markdown 必须使用 UTF-8。") from exc
            return "MARKDOWN"
        raise fail("UNSUPPORTED_FORMAT", "仅支持 PDF、PPTX、DOCX 和 Markdown。")

    def parse(self, source_type, content, document_id):
        return getattr(self, f"parse_{source_type.lower()}")(content, document_id)

    async def parse_async(self, source_type, content, document_id):
        env = {"PATH": os.defpath, "PYTHONPATH": str(Path(__file__).resolve().parents[2]),
               "PYTHONIOENCODING": "utf-8"}
        env.update({k: os.environ[k] for k in ("SYSTEMROOT", "WINDIR", "TEMP", "TMP") if k in os.environ})
        process = await asyncio.create_subprocess_exec(sys.executable, "-m",
            "concept_to_code_learning.documents.worker", source_type, document_id,
            stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.DEVNULL, env=env)
        async def send():
            process.stdin.write(content)
            await process.stdin.drain()
            process.stdin.close()
        async def read():
            chunks, size = [], 0
            while chunk := await process.stdout.read(65536):
                size += len(chunk)
                if size > 40 * 1024 * 1024:
                    raise fail("FILE_TOO_LARGE", "提取结果超过安全上限。", 413)
                chunks.append(chunk)
            return b"".join(chunks)
        pumps = [asyncio.create_task(send()), asyncio.create_task(read())]
        try:
            async with asyncio.timeout(35):
                _, payload = await asyncio.gather(*pumps)
                await process.wait()
            if process.returncode:
                raise fail("INVALID_FILE", "解析进程未完成；文件可能过于复杂或已损坏。")
            value = json.loads(payload)
            if "error" in value:
                raise fail(value["error"], value["message"], value["status"])
            return ([DocumentUnit.model_validate(u) for u in value["units"]],
                    {key: DocumentAsset(base64.b64decode(a["content"]), a["media_type"], a["file_name"])
                     for key, a in value["assets"].items()}, value["warnings"], value["capabilities"])
        except TimeoutError:
            raise fail("PROVIDER_TIMEOUT", "文档解析超时，请拆分文件后重试。", 504) from None
        finally:
            async def cleanup():
                for pump in pumps:
                    if not pump.done():
                        pump.cancel()
                if process.returncode is None:
                    try:
                        process.kill()
                    except ProcessLookupError:
                        pass
                await asyncio.gather(*pumps, return_exceptions=True)
                await process.wait()
            await settle(asyncio.create_task(cleanup()))

    def parse_pdf(self, content, document_id):
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise fail("DEPENDENCY_MISSING", "未安装 PDF 解析依赖。", 503) from exc
        with io.BytesIO(content) as stream:
            pdf = PdfReader(stream, strict=True)
            if pdf.is_encrypted:
                raise fail("INVALID_FILE", "不支持加密 PDF。")
            if not 1 <= len(pdf.pages) <= MAX_UNITS:
                raise fail("FILE_TOO_LARGE", "PDF 需包含 1 到 300 个物理页面。", 413)
            units = []
            for index, page in enumerate(pdf.pages, 1):
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 8*1024*1024:
                    raise fail("FILE_TOO_LARGE", "PDF 页面解码内容超过安全上限。", 413)
                text = (page.extract_text() or "").strip()
                blocks = [block("page", index, 1, "paragraph", text, [])] if text else []
                units.append(DocumentUnit(mode="LIVE", document_id=document_id, document_revision=1,
                    unit_id=f"page-{index}", unit_type="page", index=index, heading_path=[], blocks=blocks,
                    source_locator=locator("page", index, []),
                    preview=Preview(kind="native_pdf", asset_id="original-pdf", fidelity="original",
                        limitations=["文字阅读顺序由 PDF 编码决定；扫描页尚不支持 OCR。"]),
                    extraction_status="READY" if blocks else "NO_EXTRACTABLE_TEXT",
                    warnings=[] if blocks else ["NO_EXTRACTABLE_TEXT"]))
        return units, {}, ["PDF_READING_ORDER_MAY_DIFFER"], ["native-pdf-preview", "physical-pages", "selectable-text"]

    def parse_pptx(self, content, document_id):
        validate_office(content, "PPTX")
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise fail("DEPENDENCY_MISSING", "未安装 PPTX 解析依赖。", 503) from exc
        deck = Presentation(io.BytesIO(content))
        if not 1 <= len(deck.slides) <= MAX_UNITS:
            raise fail("FILE_TOO_LARGE", "PPTX 需包含 1 到 300 页。", 413)
        units, assets, any_unsupported = [], {}, False
        for index, slide in enumerate(deck.slides, 1):
            blocks, headings, unsupported = [], [], False
            for shape in slide.shapes:
                number = len(blocks) + 1
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    blocks.append(block("slide", index, number, "table",
                                        "\n".join(" | ".join(row) for row in rows), headings, rows))
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    if len(assets) >= MAX_ASSETS:
                        raise fail("INVALID_FILE", "PPTX 图片数量超过安全上限。")
                    if shape.image.content_type in SAFE_IMAGES:
                        asset_id = f"slide-{index}-image-{len(assets) + 1}"
                        assets[asset_id] = DocumentAsset(shape.image.blob, shape.image.content_type,
                                                         f"{asset_id}.{shape.image.ext}")
                        blocks.append(block("slide", index, number, "image", "", headings,
                                            asset_id=asset_id))
                elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                    text = shape.text.strip()
                    if shape == slide.shapes.title:
                        headings = [text]
                    blocks.append(block("slide", index, number, "paragraph", text, headings))
                else:
                    unsupported = any_unsupported = True
                    blocks.append(block("slide", index, number, "unsupported", "", headings))
            has_text = any(item.text for item in blocks)
            units.append(DocumentUnit(
                mode="LIVE", document_id=document_id, document_revision=1,
                unit_id=f"slide-{index}", unit_type="slide", index=index,
                heading_path=headings, blocks=blocks, source_locator=locator("slide", index, headings),
                preview=Preview(kind="learning_view", fidelity="partial",
                                limitations=["学习视图不保留动画、图表和原版排版。"]),
                extraction_status="PARTIAL" if unsupported else (
                    "READY" if has_text else "NO_EXTRACTABLE_TEXT"),
                warnings=["UNSUPPORTED_OFFICE_CONTENT"] if unsupported else []))
        warnings = ["OFFICE_LEARNING_VIEW_DIFFERS_FROM_ORIGINAL"]
        if any_unsupported:
            warnings.append("UNSUPPORTED_OFFICE_CONTENT")
        return units, assets, warnings, ["slides", "text-boxes", "tables", "images"]

    def parse_docx(self, content, document_id):
        validate_office(content, "DOCX")
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise fail("DEPENDENCY_MISSING", "未安装 DOCX 解析依赖。", 503) from exc
        doc = Document(io.BytesIO(content))
        assets, image_refs = {}, {}
        for rel_id, rel in doc.part.rels.items():
            if "image" in rel.reltype and not getattr(rel, "is_external", False):
                media = rel.target_part.content_type
                if media in SAFE_IMAGES:
                    if len(assets) >= MAX_ASSETS:
                        raise fail("INVALID_FILE", "DOCX 图片数量超过安全上限。")
                    asset_id = f"docx-image-{len(assets) + 1}"
                    assets[asset_id] = DocumentAsset(rel.target_part.blob, media, asset_id)
                    image_refs[rel_id] = asset_id
        units, headings, current = [], [], []
        def flush():
            if not current:
                return
            index = len(units) + 1
            values = [block("section", index, i, item.kind, item.text, headings,
                            item.table_rows, item.image_asset_id) for i, item in enumerate(current, 1)]
            units.append(DocumentUnit(
                mode="LIVE", document_id=document_id, document_revision=1,
                unit_id=f"section-{index}", unit_type="section", index=index,
                heading_path=list(headings), blocks=values,
                source_locator=locator("section", index, headings),
                preview=Preview(kind="learning_view", fidelity="partial",
                                limitations=["Word 学习视图按标题定位，不提供物理页码。"]),
                extraction_status="PARTIAL" if any(x.kind == "unsupported" for x in values) else (
                    "READY" if any(x.text for x in values) else "NO_EXTRACTABLE_TEXT")))
            current.clear()
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                match = re.match(r"Heading\s+(\d+)", item.style.name if item.style else "", re.I)
                if match and text:
                    flush()
                    level = int(match.group(1))
                    headings[:] = headings[:level - 1] + [text]
                if text:
                    current.append(block("section", 1, len(current) + 1, "paragraph", text, headings))
                for node in item._p.iter():
                    if node.tag.endswith("}blip"):
                        relation = node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                        if relation in image_refs:
                            current.append(block("section", 1, len(current) + 1, "image", "", headings,
                                                 asset_id=image_refs[relation]))
                    elif node.tag.endswith(("}oMath", "}chart")):
                        current.append(block("section", 1, len(current) + 1, "unsupported", "", headings))
            elif isinstance(item, Table):
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                current.append(block("section", 1, len(current) + 1, "table",
                                     "\n".join(" | ".join(row) for row in rows), headings, rows))
        flush()
        if not units:
            units.append(DocumentUnit(
                mode="LIVE", document_id=document_id, document_revision=1,
                unit_id="section-1", unit_type="section", index=1, heading_path=list(headings),
                blocks=[], source_locator=locator("section", 1, headings),
                preview=Preview(kind="learning_view", fidelity="partial",
                                limitations=["Word 学习视图按标题定位，不提供物理页码。"]),
                extraction_status="NO_EXTRACTABLE_TEXT"))
        return units, assets, ["OFFICE_LEARNING_VIEW_DIFFERS_FROM_ORIGINAL"], [
            "heading-paths", "paragraph-ids", "tables", "images"]

    def parse_markdown(self, content, document_id):
        text = content.decode("utf-8-sig")
        units, headings, current, code_lines, in_code = [], [], [], [], False
        def flush():
            if not current:
                return
            index = len(units) + 1
            values = [block("section", index, i, x.kind, x.text, headings)
                      for i, x in enumerate(current, 1)]
            units.append(DocumentUnit(
                mode="LIVE", document_id=document_id, document_revision=1,
                unit_id=f"section-{index}", unit_type="section", index=index,
                heading_path=list(headings), blocks=values,
                source_locator=locator("section", index, headings),
                preview=Preview(kind="learning_view", fidelity="extracted",
                                limitations=["HTML、脚本和危险链接不执行。"]),
                extraction_status="READY" if values else "NO_EXTRACTABLE_TEXT"))
            current.clear()
        for raw in text.splitlines():
            if raw.startswith("```"):
                if in_code:
                    current.append(block("section", 1, len(current) + 1, "code",
                                         "\n".join(code_lines), headings))
                    code_lines = []
                in_code = not in_code
                continue
            if in_code:
                code_lines.append(raw)
                continue
            heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", raw)
            if heading:
                flush()
                level = len(heading.group(1))
                headings[:] = headings[:level - 1] + [heading.group(2)]
                current.append(block("section", 1, len(current) + 1, "paragraph", heading.group(2), headings))
                continue
            clean = re.sub(r"<[^>]*>", "", raw).strip()
            clean = re.sub(r"\[([^]]+)\]\((?:javascript|data|file):.*\)", r"\1", clean,
                           flags=re.I)
            if clean:
                current.append(block("section", 1, len(current) + 1, "paragraph", clean, headings))
        if in_code:
            current.append(block("section", 1, len(current) + 1, "code",
                                 "\n".join(code_lines), headings))
        flush()
        if not units:
            units.append(DocumentUnit(
                mode="LIVE", document_id=document_id, document_revision=1,
                unit_id="section-1", unit_type="section", index=1, heading_path=list(headings),
                blocks=[], source_locator=locator("section", 1, headings),
                preview=Preview(kind="learning_view", fidelity="extracted",
                                limitations=["HTML、脚本和危险链接不执行。"]),
                extraction_status="NO_EXTRACTABLE_TEXT"))
        return units, {}, [], ["safe-text", "code-blocks", "heading-paths"]


def build_provider(settings):
    return FullDocumentProvider(settings)
