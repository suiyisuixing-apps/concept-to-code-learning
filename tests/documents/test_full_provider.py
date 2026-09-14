import io
from pathlib import Path

import pytest

from concept_to_code_learning.documents.full import build_provider
from concept_to_code_learning.full_contracts.models import (
    ContextRequest,
    SelectionLocator,
    SelectionSpan,
    digest,
)
from concept_to_code_learning.full_learning.context import resolve_context
from concept_to_code_learning.full_learning.errors import LearningError
from concept_to_code_learning.full_learning.ports import DocumentUpload
from concept_to_code_learning.full_learning.providers import ProviderSettings


@pytest.fixture
def provider(tmp_path):
    return build_provider(ProviderSettings(root=Path.cwd(), data_dir=tmp_path))


async def test_markdown_is_safe_unicode_persistent_and_selectable(provider):
    upload = DocumentUpload("中文 课件.md", "# 标题\nA😀中文\n<script>alert(1)</script>\n[x](javascript:evil)\n```py\nprint('好')\n```".encode())
    record = await provider.import_document(upload)
    units = await provider.list_units(record.document_id)
    assert record.source_type == "MARKDOWN" and record.unit_count == 1
    assert "<script>" not in "\n".join(x.text for x in units[0].blocks)
    assert next(x.text for x in units[0].blocks if x.text.startswith("x")) == "x"
    text_block = next(b for b in units[0].blocks if "A😀" in b.text)
    selected = "😀中"
    request = ContextRequest(document_revision=1, unit_id=units[0].unit_id,
        selected_text=selected, selected_text_hash=digest(selected),
        selection_locator=SelectionLocator(spans=[SelectionSpan(block_id=text_block.block_id, start=1, end=3)]))
    context = resolve_context(record, units[0], request)
    assert context.selected_text == selected
    assert (await provider.list_documents())[0].file_name == "中文 课件.md"


async def test_selection_mismatch_is_rejected(provider):
    record = await provider.import_document(DocumentUpload("x.md", b"hello"))
    unit = (await provider.list_units(record.document_id))[0]
    request = ContextRequest(document_revision=1, unit_id=unit.unit_id, selected_text="other",
        selected_text_hash=digest("other"), selection_locator=SelectionLocator(
            spans=[SelectionSpan(block_id=unit.blocks[0].block_id, start=0, end=5)]))
    with pytest.raises(LearningError, match="SELECTION_MISMATCH"):
        resolve_context(record, unit, request)


async def test_wrong_magic_and_corrupt_office_are_rejected_without_staging(provider):
    with pytest.raises(LearningError, match="INVALID_FILE"):
        await provider.import_document(DocumentUpload("fake.pdf", b"not pdf"))
    with pytest.raises(LearningError, match="INVALID_FILE"):
        await provider.import_document(DocumentUpload("fake.pptx", b"PK broken"))
    assert not list(provider.root.glob(".import-*"))


async def test_real_pdf_pages_include_blank_page_and_native_asset(provider):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Hello PDF) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    content = output.getvalue()
    record = await provider.import_document(DocumentUpload("pages.pdf", content))
    units = await provider.list_units(record.document_id)
    assert [x.extraction_status for x in units] == ["READY", "NO_EXTRACTABLE_TEXT"]
    asset = await provider.get_asset(record.document_id, "original-pdf")
    assert asset.content == content and asset.media_type == "application/pdf"


async def test_real_pptx_preserves_slide_table_and_picture(provider):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "中文 😀"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "键", "值"
    image = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(image, "PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(1), Inches(3))
    output = io.BytesIO()
    deck.save(output)
    record = await provider.import_document(DocumentUpload("slides.pptx", output.getvalue()))
    unit = (await provider.list_units(record.document_id))[0]
    assert unit.heading_path == ["中文 😀"]
    assert any(x.kind == "table" and x.table_rows[0] == ["键", "值"] for x in unit.blocks)
    image_block = next(x for x in unit.blocks if x.image_asset_id)
    assert (await provider.get_asset(record.document_id, image_block.image_asset_id)).media_type == "image/png"


async def test_real_docx_uses_sections_not_fake_pages(provider):
    from docx import Document
    doc = Document()
    doc.add_heading("第一章", 1)
    doc.add_paragraph("段落 😀")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "甲", "乙"
    doc.add_heading("第二章", 1)
    doc.add_paragraph("结尾")
    output = io.BytesIO()
    doc.save(output)
    record = await provider.import_document(DocumentUpload("文档.docx", output.getvalue()))
    units = await provider.list_units(record.document_id)
    assert record.unit_count == 2 and all(x.unit_type == "section" for x in units)
    assert units[0].heading_path == ["第一章"] and any(x.kind == "table" for x in units[0].blocks)
    assert "物理页码" in units[0].preview.limitations[0]
