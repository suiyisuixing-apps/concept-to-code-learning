import asyncio
import io
import re
import threading
import zipfile

import pytest

from concept_to_code_learning.documents import full
from concept_to_code_learning.full_contracts import models as m
from concept_to_code_learning.full_learning.context import resolve_context
from concept_to_code_learning.full_learning.errors import LearningError
from concept_to_code_learning.full_learning.ports import DocumentUpload
from concept_to_code_learning.full_learning.providers import ProviderSettings


@pytest.fixture
def provider(tmp_path):
    return full.build_provider(ProviderSettings(root=tmp_path, data_dir=tmp_path))


def office_package(kind, encoding, part, forbidden):
    output = io.BytesIO()
    if kind == "DOCX":
        from docx import Document
        doc = Document()
        doc.add_paragraph("正常中文 😀")
        doc.save(output)
        main = "word/document.xml"
    else:
        from pptx import Presentation
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[5])
        slide.shapes.title.text = "正常中文 😀"
        doc.save(output)
        main = "ppt/presentation.xml"
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source:
        members = {n: source.read(n) for n in source.namelist()}
    if part == "main":
        name = main
        xml = re.sub(r"^<\?xml[^>]*\?>", "", members[name].decode("utf-8")).strip()
    else:
        name = "custom/payload." + part
        xml = "<root>合法的附加 XML</root>"
        types = members["[Content_Types].xml"].decode("utf-8")
        types = types.replace("</Types>", f'<Override PartName="/{name}" ContentType="application/xml"/></Types>')
        members["[Content_Types].xml"] = types.encode("utf-8")
    doctype = '<!DOCTYPE root [<!ENTITY sample SYSTEM "file:///must-not-be-read">]>' if forbidden else ""
    members[name] = (f'<?xml version="1.0" encoding="{encoding}"?>' + doctype + xml).encode(encoding)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, raw in members.items():
            archive.writestr(name, raw)
    return output.getvalue()


@pytest.mark.parametrize("kind", ["DOCX", "PPTX"])
@pytest.mark.parametrize("encoding,part", [("utf-8", "main"), ("utf-16", "main"), ("utf-16", "bin"), ("utf-8", "XML")])
async def test_office_entity_guard_has_valid_import_control(provider, kind, encoding, part):
    valid = office_package(kind, encoding, part, False)
    record = await provider.import_document(DocumentUpload("valid." + kind.lower(), valid))
    assert record.unit_count >= 1
    poisoned = office_package(kind, encoding, part, True)
    # Check the actual guard, so a later generic parser failure cannot pass this test.
    with pytest.raises(LearningError) as rejected:
        full.validate_office(poisoned, kind)
    assert rejected.value.code == "INVALID_FILE"
    assert "XML 实体" in rejected.value.message


@pytest.mark.parametrize("part_name", ["payload.bin", "payload%20name.bin"])
async def test_referenced_xml_part_matches_downstream_opc_override_dispatch(provider, part_name):
    with zipfile.ZipFile(io.BytesIO(office_package("DOCX", "utf-8", "main", False))) as source:
        members = {name: source.read(name) for name in source.namelist()}
    path = "word/" + part_name
    members[path] = members["word/document.xml"]
    members["_rels/.rels"] = members["_rels/.rels"].replace(b'word/document.xml', path.encode())
    types = members["[Content_Types].xml"].decode("utf-8")
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
    members["[Content_Types].xml"] = types.replace("</Types>", f'<Override PartName="/{path.upper()}" ContentType="{mime}"/></Types>').encode()

    def package():
        result = io.BytesIO()
        with zipfile.ZipFile(result, "w") as archive:
            for name, raw in members.items():
                archive.writestr(name, raw)
        return result.getvalue()

    record = await provider.import_document(DocumentUpload(part_name + ".docx", package()))
    unit = (await provider.list_units(record.document_id))[0]
    assert "正常中文" in " ".join(b.text for b in unit.blocks)
    members[path] = members[path].replace(b'?>', b'?><!DOCTYPE document [<!ENTITY probe "blocked">]>', 1)
    with pytest.raises(LearningError) as rejected:
        full.validate_office(package(), "DOCX")
    assert "XML 实体" in rejected.value.message


@pytest.mark.parametrize("spoof", [
    '<x:Override xmlns:x="urn:ignored" PartName="/custom/payload.bin" ContentType="application/octet-stream"/>',
    '<ignored><Override PartName="/custom/payload.bin" ContentType="application/octet-stream"/></ignored>',
])
def test_only_direct_opc_content_type_nodes_control_xml_dispatch(spoof):
    with zipfile.ZipFile(io.BytesIO(office_package("DOCX", "utf-16", "bin", True))) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(b"</Types>", spoof.encode() + b"</Types>")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, raw in members.items():
            archive.writestr(name, raw)
    with pytest.raises(LearningError) as rejected:
        full.validate_office(output.getvalue(), "DOCX")
    assert "XML 实体" in rejected.value.message


def test_content_type_extension_uses_lower_without_casefold_collisions():
    with zipfile.ZipFile(io.BytesIO(office_package("DOCX", "utf-16", "bin", True))) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["custom/payload.ß"] = members.pop("custom/payload.bin")
    types = members["[Content_Types].xml"].decode("utf-8")
    types = re.sub(r'<Override PartName="/custom/payload.bin"[^>]*/>', "", types)
    types = types.replace("</Types>", '<Default Extension="ß" ContentType="application/xml"/><Default Extension="ss" ContentType="application/octet-stream"/></Types>')
    members["[Content_Types].xml"] = types.encode()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, raw in members.items():
            archive.writestr(name, raw)
    with pytest.raises(LearningError) as rejected:
        full.validate_office(output.getvalue(), "DOCX")
    assert "XML 实体" in rejected.value.message


async def test_metadata_limit_rejects_before_publication(provider, monkeypatch):
    monkeypatch.setattr(full, "MAX_METADATA_BYTES", 200)
    with pytest.raises(LearningError, match="FILE_TOO_LARGE"):
        await provider.import_document(DocumentUpload("oversize.md", b"# heading\nbody"))
    assert list(provider.root.iterdir()) == []


async def test_bom_and_long_context_preserve_original_selection(provider):
    content = "# 标题\n" + "\n\n".join(["甲" * 40000] * 3) + "😀选区"
    record = await provider.import_document(DocumentUpload("long.md", content.encode("utf-8-sig")))
    unit = (await provider.list_units(record.document_id))[0]
    assert unit.heading_path == ["标题"]
    block = next(b for b in unit.blocks if b.text.endswith("😀选区"))
    request = m.ContextRequest(document_revision=1, unit_id=unit.unit_id,
        selected_text="😀选区", selected_text_hash=m.digest("😀选区"),
        selection_locator=m.SelectionLocator(spans=[m.SelectionSpan(
            block_id=block.block_id, start=len(block.text) - 3, end=len(block.text))]))
    context = resolve_context(record, unit, request)
    assert len(context.visible_text) == 100000
    assert context.coverage == "PARTIAL" and "CONTEXT_OVERVIEW_TRUNCATED" in context.warnings
    assert context.selected_text == "😀选区"
    assert next(b for b in context.relevant_context_blocks if b.block_id == block.block_id).text == block.text


async def test_slow_document_hash_does_not_block_event_loop(provider, monkeypatch):
    record = await provider.import_document(DocumentUpload("normal.md", b"hello"))
    entered, release = threading.Event(), threading.Event()
    real_record = provider.record

    def slow_record(folder):
        entered.set()
        assert release.wait(3)
        return real_record(folder)

    monkeypatch.setattr(provider, "record", slow_record)
    pending = asyncio.create_task(provider.get_document(record.document_id))
    try:
        assert await asyncio.to_thread(entered.wait, 2)
        assert not pending.done()  # This coroutine remains responsive during file IO.
        await asyncio.sleep(0)
    finally:
        release.set()
    assert (await pending).document_id == record.document_id


async def test_cancelled_import_never_publishes_staged_document(provider, monkeypatch):
    original = provider._stage
    entered, release = threading.Event(), threading.Event()

    def held_stage(*args):
        result = original(*args)
        entered.set()
        assert release.wait(3)
        return result

    monkeypatch.setattr(provider, "_stage", held_stage)
    pending = asyncio.create_task(provider.import_document(DocumentUpload("cancel.md", b"hello")))
    try:
        assert await asyncio.to_thread(entered.wait, 2)
        pending.cancel()
        await asyncio.sleep(0)
        assert not pending.done()
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError):
        await pending
    assert list(provider.root.iterdir()) == []
